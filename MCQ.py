import streamlit as st
import re
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from io import BytesIO
from ai import ask_ai
from concurrent.futures import ThreadPoolExecutor, as_completed


# ---------------- GENERATE ONE MCQ ----------------
def generate_single_mcq(topic, difficulty, q_number):

    system_prompt = """
You are an expert educational assessment designer.
Generate high-quality multiple-choice questions.
"""

    user_prompt = f"""
Generate ONE multiple-choice question in Traditional Chinese (Taiwan).

Topic: {topic}
Difficulty: {difficulty}

Rules:
- 4 options only (a,b,c,d)
- Options must be short; question and options in Traditional Chinese
- Include correct answer
- No coding questions

Format EXACTLY:

Q{q_number}: Question text
a. Option
b. Option
c. Option
d. Option
Answer: a
"""

    return ask_ai(system_prompt, user_prompt)


# ---------------- FORMAT QUIZ ----------------
def format_quiz(quiz_text):

    lines = quiz_text.split("\n")

    formatted_quiz = []
    current_question = []

    for line in lines:

        if re.match(r"^Q\d+:", line):
            if current_question:
                formatted_quiz.append(current_question)
            current_question = [line]

        elif re.match(r"^[a-d]\.", line):
            current_question.append(line)

        elif line.startswith("Answer:"):
            current_question.append(line)

    if current_question:
        formatted_quiz.append(current_question)

    return formatted_quiz


# ---------------- BALANCED DIFFICULTY ----------------
def create_difficulty_mix(total):

    beginner = total // 3
    intermediate = total // 3
    expert = total - beginner - intermediate

    difficulties = (
        ["Beginner"] * beginner
        + ["Intermediate"] * intermediate
        + ["Expert"] * expert
    )

    return difficulties


# ---------------- PARALLEL QUIZ GENERATION ----------------
def generate_quiz_parallel(topic, difficulty, num_questions, balanced):

    results = []

    if balanced:
        difficulty_list = create_difficulty_mix(num_questions)
    else:
        difficulty_list = [difficulty] * num_questions

    with ThreadPoolExecutor(max_workers=5) as executor:

        futures = []

        for i, diff in enumerate(difficulty_list, start=1):
            futures.append(
                executor.submit(generate_single_mcq, topic, diff, i)
            )

        for future in as_completed(futures):
            results.append(future.result())

    return results


# ---------------- CREATE DOCX ----------------
def generate_docx(quiz, heading1, heading2):

    doc = Document()

    h1 = doc.add_heading(heading1, level=0)
    h1.alignment = WD_ALIGN_PARAGRAPH.CENTER

    h2 = doc.add_heading(heading2, level=2)
    h2.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph("姓名：")
    doc.add_paragraph("座號：")
    doc.add_paragraph("班級：")
    doc.add_paragraph("組別：")
    doc.add_paragraph("")

    # Questions
    for question in quiz:

        for line in question:
            if not line.startswith("Answer:"):
                doc.add_paragraph(line)

        doc.add_paragraph("")

    # Answer Key
    doc.add_page_break()
    doc.add_heading("參考答案", level=1)

    for i, question in enumerate(quiz):

        for line in question:
            if line.startswith("Answer:"):
                ans = line.split(": ")[1]
                doc.add_paragraph(f"Q{i+1}: {ans}")

    buffer = BytesIO()
    doc.save(buffer)
    buffer.seek(0)

    return buffer


# ---------------- STREAMLIT APP ----------------
def MCQ():

    st.subheader("AI 出題產生器")

    col1, col2 = st.columns(2)

    with col1:
        institute_name = st.text_input("學校／補習班名稱")
        topic = st.text_input("出題範圍（例：國一 一元一次方程式）")

    with col2:
        quiz_title = st.text_input("考卷標題")

        DIFF_ZH = {"Beginner": "初級", "Intermediate": "中級", "Expert": "進階"}
        difficulty = st.selectbox(
            "難易度",
            ["Beginner", "Intermediate", "Expert"],
            format_func=lambda k: DIFF_ZH.get(k, k),
        )

    num_questions = st.number_input(
        "題數",
        min_value=1,
        max_value=20,
        value=5
    )

    balanced = st.checkbox(
        "難易度混搭（初／中／高自動分配）",
        value=False
    )

    show_answers = st.checkbox("顯示答案", value=True)

    # ---------- GENERATE QUIZ ----------
    if st.button("開始出題"):

        if topic.strip() == "":
            st.error("請先輸入出題範圍。")
            return

        with st.spinner("AI 出題中……"):

            raw_questions = generate_quiz_parallel(
                topic,
                difficulty,
                num_questions,
                balanced
            )

        formatted_quiz = []

        for q in raw_questions:
            formatted_quiz.extend(format_quiz(q))

        if not formatted_quiz:
            st.error("AI 沒有回傳可用題目，原始回應如下（拿去對金鑰或額度）：")
            for q in raw_questions:
                st.code(q or "(空白回應)")
            return

        st.session_state["quiz"] = formatted_quiz

        docx_content = generate_docx(
            formatted_quiz,
            institute_name or "Institute",
            quiz_title or "Quiz"
        )

        st.session_state["docx_content"] = docx_content

    # ---------- DISPLAY QUIZ ----------
    if "quiz" in st.session_state:

        st.divider()

        if institute_name:
            st.write(f"### {institute_name}")

        if quiz_title:
            st.write(f"**{quiz_title}**")

        st.write("---")

        st.write("姓名：")
        st.write("座號：")
        st.write("班級：")
        st.write("組別：")
        st.write("")

        for question in st.session_state["quiz"]:

            st.markdown(f"**{question[0]}**")

            for line in question[1:]:

                if line.startswith("Answer:"):

                    if show_answers:
                        st.success(f"正解：{line.split(': ')[1]}")

                else:
                    st.write(line)

            st.write("")

        # ---------- ANSWER KEY ----------
        if show_answers:

            st.subheader("參考答案")

            for i, question in enumerate(st.session_state["quiz"]):

                for line in question:
                    if line.startswith("Answer:"):
                        ans = line.split(": ")[1]
                        st.write(f"Q{i+1}: {ans}")

        # ---------- DOWNLOAD ----------
        if "docx_content" in st.session_state:

            st.download_button(
                label="下載考卷（DOCX）",
                data=st.session_state["docx_content"],
                file_name="quiz.docx",
                mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )