import streamlit as st
from ai import ask_ai
from docx import Document
from io import BytesIO


# ---------------- CREATE DOCX ----------------

def create_docx(chat_history):

    doc = Document()
    doc.add_heading("心情樹洞對話紀錄", level=1)

    for msg in chat_history:
        role = msg["role"].capitalize()
        content = msg["content"]
        doc.add_paragraph(f"{role}: {content}")

    buffer = BytesIO()
    doc.save(buffer)
    buffer.seek(0)

    return buffer


# ---------------- COUNSELLOR ----------------

def counsellor():

    st.subheader("🧠 教師心情樹洞")

    st.markdown(
        "聊聊備課壓力、職業倦怠或班級經營的煩惱。對話可下載留存。"
    )

    st.info(
        "⚠️ 這是支持性談話，不是專業心理諮商；有需要請尋求專業協助。"
    )

    # ---------------- RESET CHAT ----------------

    if st.button("開始新對話"):
        st.session_state.messages = []
        st.rerun()


    # ---------------- QUICK TOPICS ----------------

    st.markdown("### 常見話題")

    col1, col2, col3 = st.columns(3)

    if col1.button("壓力大"):
        st.session_state.messages.append(
            {"role": "user", "content": "我最近備課量很大，覺得壓力很大。"}
        )
        st.rerun()

    if col2.button("倦怠感"):
        st.session_state.messages.append(
            {"role": "user", "content": "教久了覺得倦怠，提不起勁。"}
        )
        st.rerun()

    if col3.button("班級經營"):
        st.session_state.messages.append(
            {"role": "user", "content": "班上有難處理的秩序問題，不知道怎麼辦。"}
        )
        st.rerun()


    # ---------------- INIT CHAT ----------------

    if "messages" not in st.session_state:
        st.session_state.messages = []


    # ---------------- DISPLAY CHAT ----------------

    for message in st.session_state.messages:

        with st.chat_message(message["role"]):
            st.markdown(message["content"])


    st.divider()


    # ---------------- USER INPUT ----------------

    if prompt := st.chat_input("今天想聊什麼？"):

        st.session_state.messages.append(
            {"role": "user", "content": prompt}
        )

        with st.chat_message("user"):
            st.markdown(prompt)


        # ---------------- AI RESPONSE ----------------

        with st.chat_message("assistant"):

            with st.spinner("想一想……"):

                MAX_HISTORY = 10

                recent_messages = st.session_state.messages[-MAX_HISTORY:]

                conversation = ""

                for msg in recent_messages:
                    role = msg["role"]
                    content = msg["content"]
                    conversation += f"{role}: {content}\n"


                system_prompt = """
You are a warm and compassionate AI counsellor supporting Taiwanese teachers. Reply in Traditional Chinese.

Speak like a supportive human conversation partner.

Guidelines:
- Start with empathy
- Use natural conversational language
- Keep responses short and friendly
- Give only 1–2 practical suggestions
- Ask thoughtful follow-up questions

Tone:
- calm
- supportive
- friendly
- conversational

Do NOT:
- give medical diagnoses
- sound like an academic essay
- give long lectures

Structure responses like a real conversation:
1. Empathy
2. Small helpful idea
3. Gentle question
"""


                user_prompt = f"""
Conversation so far:
{conversation}

Respond to the latest teacher message in a natural, human-like way.
"""


                response = ask_ai(system_prompt, user_prompt)

            st.markdown(response)


        st.session_state.messages.append(
            {"role": "assistant", "content": response}
        )


    # ---------------- DOWNLOAD DOCX ----------------

    if st.session_state.messages:

        docx_file = create_docx(st.session_state.messages)

        st.download_button(
            label="⬇ 下載對話（DOCX）",
            data=docx_file,
            file_name="teacher_wellness_chat.docx",
            mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        )