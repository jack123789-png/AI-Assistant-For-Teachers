import streamlit as st
from ai import ask_ai
from docx import Document
from io import BytesIO
from pptx import Presentation


# ---------------- CLEAN DISPLAY ----------------
def clean_lesson_plan_display(text):

    lines = text.split("\n")
    cleaned = []

    for line in lines:
        if "-------------------------------------" in line:
            cleaned.append("")
        else:
            cleaned.append(line)

    return "\n".join(cleaned)


# ---------------- GENERATE LESSON PLAN ----------------
def generate_lesson_plan(country, class_level, subject, topic, duration, sessions):

    system_prompt = """
You are an expert teacher trainer familiar with Taiwan 12-year Basic Education Curriculum Guidelines (108課綱).

Generate professional lesson plans for Taiwanese elementary and junior high teachers.
"""

    user_prompt = f"""
Country: 台灣
課綱：十二年國教108課綱
Class Level: {class_level}
Subject: {subject}

Topic:
{topic}

每節時間：{duration} 分鐘
Number of Sessions: {sessions}

Create a structured professional lesson plan.

Use this structure:

-------------------------------------
LESSON IDENTIFICATION
School:
Teacher:
Subject:
Class:
Duration:
Topic:
Sub-topic:

-------------------------------------
COMPETENCY
Key competency to be developed

-------------------------------------
LEARNING OBJECTIVES
• objective
• objective

-------------------------------------
KEY INQUIRY QUESTION

-------------------------------------
LEARNING ACTIVITIES
Teacher Activities
Learner Activities

-------------------------------------
ASSESSMENT METHODS

-------------------------------------
LEARNING RESOURCES

-------------------------------------
TEACHER REFLECTION

-------------------------------------

Ensure explanations match the cognitive level of {class_level} learners in Taiwan. Use Traditional Chinese.
"""

    return ask_ai(system_prompt, user_prompt)


# ---------------- GENERATE SLIDE CONTENT ----------------
def generate_slide_content(country, class_level, subject, topic):

    system_prompt = """
You are a professional teacher and presentation designer.
Generate teaching slides appropriate for classroom instruction.
"""

    user_prompt = f"""
Country: {country}
Class Level: {class_level}
Subject: {subject}

Topic:
{topic}

Create structured teaching slides.

Rules:
• Maximum 5 bullet points
• Clear titles
• Professional explanations
• Include speaker notes

Use the format:

Slide 1
Title: ...
Content:
- point
- point
Notes:
teacher explanation

Slide 2
Title: ...
Content:
- point
- point
Notes:
teacher explanation
"""

    return ask_ai(system_prompt, user_prompt)


# ---------------- CREATE DOCX ----------------
def generate_docx(text):

    doc = Document()
    doc.add_heading("教案", level=0)

    lines = text.split("\n")

    for line in lines:

        if "-------------------------------------" in line:
            doc.add_paragraph("")
        else:
            doc.add_paragraph(line)

    buffer = BytesIO()
    doc.save(buffer)
    buffer.seek(0)

    return buffer


# ---------------- CREATE PPT ----------------
def generate_ppt(slide_text, topic):

    prs = Presentation()

    slide_layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(slide_layout)

    slide.shapes.title.text = topic
    slide.placeholders[1].text = "上課講義"

    slides = slide_text.split("Slide")

    for block in slides:

        if "Title:" not in block:
            continue

        lines = block.split("\n")

        title = ""
        content = []
        notes = ""
        mode = None

        for line in lines:

            if line.startswith("Title:"):
                title = line.replace("Title:", "").strip()

            elif line.startswith("Content"):
                mode = "content"

            elif line.startswith("Notes"):
                mode = "notes"

            elif line.startswith("-") and mode == "content":
                content.append(line.replace("-", "").strip())

            elif mode == "notes":
                notes += line + " "

        slide_layout = prs.slide_layouts[1]
        slide_obj = prs.slides.add_slide(slide_layout)

        slide_obj.shapes.title.text = title

        tf = slide_obj.placeholders[1].text_frame
        tf.clear()

        for point in content:
            p = tf.add_paragraph()
            p.text = point
            p.level = 0

        slide_obj.notes_slide.notes_text_frame.text = notes

    buffer = BytesIO()
    prs.save(buffer)
    buffer.seek(0)

    return buffer


# ---------------- STREAMLIT UI ----------------
def lessonplan():

    st.subheader("AI 教案產生器")

    country = "台灣"

    st.divider()

    education_level = st.selectbox(
        "學制",
        ["國小", "國中"]
    )

    if education_level == "國小":

        class_level = st.selectbox(
            "年級",
            ["國小1年級","國小2年級","國小3年級","國小4年級","國小5年級","國小6年級"]
        )

        subject = st.selectbox(
            "科目",
            ["國語","數學","英語","自然科學","社會","綜合活動","生活","健康與體育","藝術"]
        )

    else:

        class_level = st.selectbox(
            "年級",
            ["國中7年級","國中8年級","國中9年級"]
        )

        subject = st.selectbox(
            "科目",
            ["國文","英語","數學","自然科學","社會","綜合活動","健康與體育","藝術","科技"]
        )


    topic = st.text_area("單元主題／內容", height=150)

    duration = st.number_input("每節分鐘數", min_value=1, value=45)
    sessions = st.number_input("節數", min_value=1)

    st.divider()

    col1, col2 = st.columns(2)

    # -------- LESSON PLAN --------
    with col1:

        st.markdown("### 教案")
        st.markdown("---")

        if st.button("產生教案"):

            with st.spinner("教案產生中……"):

                lesson_plan = generate_lesson_plan(
                    country,
                    class_level,
                    subject,
                    topic,
                    duration,
                    sessions
                )

            st.session_state.lesson_plan = lesson_plan

        if "lesson_plan" in st.session_state:

            cleaned_plan = clean_lesson_plan_display(st.session_state.lesson_plan)

            st.text(cleaned_plan)

            docx = generate_docx(st.session_state.lesson_plan)

            st.download_button(
                "下載教案（DOCX）",
                data=docx,
                file_name="lesson_plan.docx",
                mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
            )

    # -------- SLIDES --------
    with col2:

        st.markdown("### 上課投影片")
        st.markdown("---")

        if st.button("產生投影片大綱"):

            with st.spinner("大綱產生中……"):

                slides = generate_slide_content(
                    country,
                    class_level,
                    subject,
                    topic
                )

            st.session_state.slides = slides

        if "slides" in st.session_state:

            st.markdown(st.session_state.slides)

            ppt = generate_ppt(st.session_state.slides, topic)

            st.download_button(
                "下載投影片（PPTX）",
                data=ppt,
                file_name="lesson_slides.pptx",
                mime="application/vnd.openxmlformats-officedocument.presentationml.presentation"
            )